#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para generar la Presentación Oral Oficial en PowerPoint (.pptx)
ajustada a los 15 minutos de exposición para el Proyecto Integrador
de Métodos Numéricos (UCC 2026-2).

Modelado Numérico del Resurgimiento de 'Rosa Pastel' (Belanova) en Spotify México (2023).
Datos reales de kworb.net / Spotify Charts México.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # carpeta del proyecto
    blank_layout = prs.slide_layouts[6]

    # Paleta de colores
    navy = RGBColor(31, 78, 121)        # #1F4E79 UCC Navy
    dark_blue = RGBColor(0, 32, 96)      # #002060
    magenta = RGBColor(233, 30, 99)      # #E91E63 Rosa Pastel Accent
    gray = RGBColor(89, 89, 89)
    dark_text = RGBColor(40, 40, 40)
    white = RGBColor(255, 255, 255)
    light_bg = RGBColor(245, 247, 250)
    border_gray = RGBColor(217, 225, 242)

    def add_header(slide, title_text, category_text="TALLER DE MÉTODOS NUMÉRICOS · PROYECTO INTEGRADOR"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = 'Calibri'
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = magenta

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = 'Calibri'
        p1.font.size = Pt(23)
        p1.font.bold = True
        p1.font.color.rgb = navy

    # -------------------------------------------------------------
    # SLIDE 1: Portada
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = dark_blue
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(11.0), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "UNIVERSIDAD CATÓLICA DE CÓRDOBA · FACULTAD DE INGENIERÍA"
    p.font.name = 'Calibri'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(189, 215, 238)
    p.space_after = Pt(20)

    p = tf1.add_paragraph()
    p.text = "Modelado Numérico del Resurgimiento de 'Rosa Pastel' en Spotify México (2023)"
    p.font.name = 'Calibri'
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = white
    p.space_after = Pt(16)

    p = tf1.add_paragraph()
    p.text = "Ajuste de Curvas (Mínimos Cuadrados), Búsqueda de Raíces (Bisección vs. Newton-Raphson) y Diferenciación Numérica"
    p.font.name = 'Calibri'
    p.font.size = Pt(15)
    p.font.italic = True
    p.font.color.rgb = RGBColor(242, 242, 242)
    p.space_after = Pt(40)

    p = tf1.add_paragraph()
    p.text = "Autores: Pedro Gatti · Facundo Grillo · Gonzalo Marin   |   Docente: Ing. Carlos Arnaudo\nIngeniería Informática   ·   Comisión 46324   ·   Octubre 2026"
    p.font.name = 'Calibri'
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(220, 220, 220)

    # -------------------------------------------------------------
    # SLIDE 2: Contexto y Justificación
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "El fenómeno de 'Rosa Pastel' y su relevancia en Ingeniería", "CONTEXTO Y JUSTIFICACIÓN")

    # Tarjeta 1: Caso Rosa Pastel
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    c1.fill.solid()
    c1.fill.fore_color.rgb = light_bg
    c1.line.color.rgb = border_gray
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_top = Inches(0.4)

    p = tf_c1.paragraphs[0]
    p.text = "El caso 'Rosa Pastel' (Belanova)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(12)

    p = tf_c1.add_paragraph()
    p.text = (
        "• Canción electropop mexicana lanzada como sencillo en 2006 (álbum Dulce Beat).\n\n"
        "• En 2023 experimentó un fuerte resurgimiento en plataformas de streaming a partir de un trend masivo en TikTok ('sueños frustrados', julio 2023).\n\n"
        "• En Spotify México las reproducciones semanales pasaron de 1.1 M a un pico de 4.17 M en julio.\n\n"
        "• El volumen total acumulado a lo largo de 2023 llegó a 132,9 millones de reproducciones en el mercado mexicano."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    # Tarjeta 2: Relevancia en Ingeniería
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1))
    c2.fill.solid()
    c2.fill.fore_color.rgb = light_bg
    c2.line.color.rgb = border_gray
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_top = Inches(0.4)

    p = tf_c2.paragraphs[0]
    p.text = "Relevancia en Ingeniería"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(12)

    p = tf_c2.add_paragraph()
    p.text = (
        "• Plataformas de streaming: Modelar cuantitativamente la evolución temporal de la demanda permite anticipar picos de consumo.\n\n"
        "• Infraestructura cloud y redes CDN: Estimar ritmos de crecimiento ayuda a dimensionar preventivamente servidores y servicios de distribución de contenido.\n\n"
        "• Naturaleza de los datos: Las mediciones de telemetría son discretas (semanales) y presentan variaciones naturales.\n\n"
        "• Rol de los métodos numéricos: Brindan el marco formal para deducir un modelo analítico continuo, localizar hitos acumulados y calcular tasas de cambio instantáneas."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    # -------------------------------------------------------------
    # SLIDE 3: Datos Experimentales (Spotify México 2023)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Datos Reales: 47 Semanas de Spotify México en 2023", "DATOS EXPERIMENTALES")

    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.1), Inches(5.1))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "Serie Temporal de Telemetría:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(10)

    p = tf3.add_paragraph()
    p.text = (
        "• Fuente oficial: kworb.net (track 6Z0ybnsUBhbZ9ZjPW64eoF), que recopila los charts semanales del Top 200 de Spotify México.\n\n"
        "• Alcance: Exclusivamente streams de Spotify México en el Top 200 semanal oficial.\n\n"
        "• Rango: 47 semanas consecutivas de 2023 (t = 0 al 26/01/2023 hasta t = 46 al 14/12/2023).\n\n"
        "• Acumulado anual: 132,9 millones de streams.\n\n"
        "• Hito de 100M: Cruzado entre la semana 35 (99.60 M) y la semana 36 (102.74 M), fines de septiembre de 2023."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    img1_path = os.path.join(base_dir, 'grafico_1_ajuste_curva_streams.png')
    if os.path.exists(img1_path):
        s3.shapes.add_picture(img1_path, Inches(6.1), Inches(1.8), width=Inches(6.4))

    # -------------------------------------------------------------
    # SLIDE 4: Método 1 - Ajuste de Curvas (Mínimos Cuadrados)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Método 1: Ajuste de Curvas por Mínimos Cuadrados", "REGRESIÓN POLINÓMICA")

    c4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
    c4.fill.solid()
    c4.fill.fore_color.rgb = light_bg
    c4.line.color.rgb = border_gray
    tf4 = c4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = Inches(0.4)

    p = tf4.paragraphs[0]
    p.text = "Modelo Continuo Cúbico: S(t) = a·t³ + b·t² + c·t + d"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(12)

    p = tf4.add_paragraph()
    p.text = (
        "• Justificación del grado 3: Permite capturar la concavidad positiva inicial (aceleración por el trend de TikTok) "
        "y la posterior concavidad negativa (inflexión y estabilización hacia fin de año).\n\n"
        "Coeficientes obtenidos (Mínimos Cuadrados sobre 47 semanas):\n"
        "   a = -0.000839  |  b = 0.076977  |  c = 1.081162  |  d = 2.072898\n\n"
        "Ecuación del modelo ajustado:\n"
        "   S(t) = -0.000839·t³ + 0.076977·t² + 1.081162·t + 2.072898\n\n"
        "Evaluación estadística de la bondad de ajuste:\n"
        "   • Suma de cuadrados de los residuos: Sr = 36.9406\n"
        "   • Suma total de cuadrados respecto a la media: St = 79,137.5830\n"
        "   • Coeficiente de determinación: R² = 0.999533 (99.95% de la varianza explicada)\n"
        "   • Error estándar de la estimación: Sy/x = 0.9269 millones de streams"
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    # -------------------------------------------------------------
    # SLIDE 5: Método 2 - Búsqueda de Raíces (Hito de 100M)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Método 2: Búsqueda de Raíces — El Hito de 100 Millones", "BÚSQUEDA DE RAÍCES")

    c5a = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    c5a.fill.solid()
    c5a.fill.fore_color.rgb = light_bg
    c5a.line.color.rgb = border_gray
    tf5a = c5a.text_frame
    tf5a.word_wrap = True
    tf5a.margin_left = tf5a.margin_top = Inches(0.4)

    p = tf5a.paragraphs[0]
    p.text = "Planteo del Problema"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(10)

    p = tf5a.add_paragraph()
    p.text = (
        "Se busca determinar la semana t* en que la canción superó 100 millones de streams en Spotify México:\n\n"
        "f(t) = S(t) - 100.0 = 0\n\n"
        "f(t) = -0.000839·t³ + 0.076977·t² + 1.081162·t - 97.927102 = 0\n\n"
        "Condición de parada exigida:\n"
        "|Ea%| < 10⁻⁴ %  (0.0001% de error relativo)\n\n"
        "Intervalo inicial [35, 36]:\n"
        "• Datos: S(35) = 99.60 M < 100 < S(36) = 102.74 M\n"
        "• Modelo: f(35) = -1.7584 < 0,  f(36) = +1.6162 > 0\n"
        "• Hay cambio de signo (Bolzano)."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    c5b = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1))
    c5b.fill.solid()
    c5b.fill.fore_color.rgb = light_bg
    c5b.line.color.rgb = border_gray
    tf5b = c5b.text_frame
    tf5b.word_wrap = True
    tf5b.margin_left = tf5b.margin_top = Inches(0.4)

    p = tf5b.paragraphs[0]
    p.text = "Resultados Comparados"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(10)

    p = tf5b.add_paragraph()
    p.text = (
        "1. Método de Bisección (Método Cerrado):\n"
        "• Intervalo inicial: [35.0, 36.0]\n"
        "• Convergencia lineal (factor 0.5 por paso).\n"
        "• 15 iteraciones requeridas.\n"
        "• Raíz hallada: t* = 35.520172 semanas.\n\n"
        "2. Método de Newton-Raphson (Método Abierto):\n"
        "• Estimación inicial: t₀ = 35.0 semanas.\n"
        "• Derivada f'(t) = -0.002517·t² + 0.153954·t + 1.081162\n"
        "• Convergencia cuadrática.\n"
        "• 3 iteraciones requeridas (5 veces más rápido).\n"
        "• Raíz hallada: t* = 35.520148 semanas.\n\n"
        "Control por interpolación lineal: t = 35.127 sem.\n"
        "Diferencia: 2.75 días (error de ajuste ±0.93 M)."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = dark_text

    # -------------------------------------------------------------
    # SLIDE 6: Comparativa Bisección vs. Newton-Raphson
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Comparativa: Bisección (Lineal) vs. Newton-Raphson (Cuadrática)", "CONVERGENCIA")

    tb6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.1), Inches(5.1))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "Análisis de Convergencia:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(10)

    p = tf6.add_paragraph()
    p.text = (
        "• Bisección: 15 iteraciones.\n"
        "  - Método cerrado, muy robusto.\n"
        "  - Incondicionalmente convergente si Bolzano se cumple.\n"
        "  - Orden lineal: divide el intervalo a la mitad en cada ciclo.\n\n"
        "• Newton-Raphson: 3 iteraciones.\n"
        "  - Método abierto por tangentes sucesivas.\n"
        "  - Orden cuadrático: duplica aproximadamente los dígitos exactos en cada paso.\n"
        "  - Requiere derivada analítica y una buena aproximación inicial.\n\n"
        "• Coincidencia en la raíz: t* ≈ 35.520 semanas.\n"
        "  - Fecha estimada: fines de septiembre / 1 de octubre de 2023."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    img2_path = os.path.join(base_dir, 'grafico_2_convergencia_biseccion_vs_newton.png')
    if os.path.exists(img2_path):
        s6.shapes.add_picture(img2_path, Inches(6.1), Inches(1.8), width=Inches(6.4))

    # -------------------------------------------------------------
    # SLIDE 7: Método 3 - Diferenciación Numérica
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Método 3: Diferenciación Numérica y Pico de Demanda", "DIFERENCIACIÓN NUMÉRICA")

    tb7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.1), Inches(5.1))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "Ritmo Semanal de Streams (dS/dt):"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(10)

    p = tf7.add_paragraph()
    p.text = (
        "• Esquemas de diferencias finitas (h = 1 semana):\n"
        "  - Adelante y atrás: error O(h).\n"
        "  - Centrada: error O(h²), mejor aproximación teórica.\n\n"
        "• Pico según el modelo (S''(t) = 0):\n"
        "  - t_pico = -b / (3a) = 30.59 semanas (fines de agosto 2023).\n"
        "  - Ritmo del modelo: 3.436 M streams/semana.\n\n"
        "• Pico real en los datos observados:\n"
        "  - Semana 26 (27/07/2023): 4.171 M streams.\n"
        "  - Coincide con el auge de TikTok de julio 2023.\n\n"
        "• Efecto de suavizado del modelo:\n"
        "  - El polinomio cúbico desplaza el pico ~4 semanas y atenúa su valor."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = dark_text

    img3_path = os.path.join(base_dir, 'grafico_3_derivada_tasa_viralidad.png')
    if os.path.exists(img3_path):
        s7.shapes.add_picture(img3_path, Inches(6.1), Inches(1.8), width=Inches(6.4))

    # -------------------------------------------------------------
    # SLIDE 8: Generalización a Otras Canciones
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Generalización: ¿sirve para cualquier canción?", "VALIDACIÓN DEL MODELO Y GENERALIZACIÓN")

    tb8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.0), Inches(5.1))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "Evaluación en 3 Canciones (Resurgimiento vs. Estrenos):"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(8)

    p = tf8.add_paragraph()
    p.text = (
        "• Regla matemática de validación:\n"
        "  - Punto de inflexión: t = -b / (3a). Es MÁXIMO del ritmo S'(t) sólo si a < 0.\n"
        "  - Si a > 0, es un MÍNIMO: la cúbica no puede representar el pico de demanda.\n\n"
        "• Resurgimiento gradual vs. Estrenos masivos:\n"
        "  - Rosa Pastel: sube meses y baja (a = -0.000839 < 0) → Aplica (pico en sem. 30.6).\n"
        "  - Estrenos (Ella Baila Sola, Tití): arrancan en pico y sólo bajan (a > 0) → No aplica.\n"
        "    La derivada parabólica vuelve a subir al final (rebote irreal en el gráfico).\n\n"
        "• Un R² alto no garantiza validez física:\n"
        "  - En los estrenos R² es casi perfecto (0.996 y 0.9997), pero el modelo falla para el pico.\n\n"
        "• Herramienta general de ingeniería reutilizable:\n"
        "  - Hito 100M (Newton en 3 iteraciones) y pico de datos funcionan para todas.\n"
        "  - El programa avisa automáticamente cuándo no aplica en lugar de arrojar un resultado falso."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = dark_text

    img4_path = os.path.join(base_dir, 'grafico_4_comparacion_canciones.png')
    if os.path.exists(img4_path):
        s8.shapes.add_picture(img4_path, Inches(8.1), Inches(1.8), height=Inches(5.1))

    # -------------------------------------------------------------
    # SLIDE 9: Limitaciones del Modelo Numérico
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Limitaciones del Modelo Numérico", "ANÁLISIS CRÍTICO")

    cards_info = [
        ("1. Discrepancia con Datos Puntuales",
         "La cúbica por mínimos cuadrados no pasa exactamente por los puntos. Por eso la raíz del modelo (t = 35.52) difiere ~2.75 días de la interpolación lineal directa entre las semanas 35 y 36 (t = 35.13).",
         Inches(0.8), Inches(1.8)),
        ("2. Suavizado del Pico de Demanda",
         "Al ajustar un único polinomio para 47 semanas, el modelo atenúa los cambios bruscos: desplaza el pico de la semana 26 a la semana 30.6 y subestima su magnitud máxima (3.44 vs. 4.17 M streams/sem).",
         Inches(6.9), Inches(1.8)),
        ("3. Alcance de los Datos Relevados",
         "Los datos cubren únicamente los charts semanales del Top 200 de Spotify México (kworb.net). No contemplan otras plataformas (YouTube, Apple Music), otros países ni escuchas fuera del chart.",
         Inches(0.8), Inches(4.5)),
        ("4. Imposibilidad de Extrapolación",
         "El modelo cúbico solo es válido dentro del rango de 2023 (t ∈ [0, 46]). Dado que su coeficiente cúbico es negativo (a = -0.000839), extrapolar fuera del año haría tender los streams acumulados a -∞.",
         Inches(6.9), Inches(4.5))
    ]

    for title_card, text_card, left_pos, top_pos in cards_info:
        card_shape = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(5.6), Inches(2.4))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = light_bg
        card_shape.line.color.rgb = border_gray
        tf_card = card_shape.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = tf_card.margin_top = Inches(0.3)

        p = tf_card.paragraphs[0]
        p.text = title_card
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = navy
        p.space_after = Pt(6)

        p = tf_card.add_paragraph()
        p.text = text_card
        p.font.size = Pt(11.5)
        p.font.color.rgb = dark_text

    # -------------------------------------------------------------
    # SLIDE 10: Conclusiones
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Conclusiones del Proyecto Integrador", "SÍNTESIS FINAL")

    c10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
    c10.fill.solid()
    c10.fill.fore_color.rgb = light_bg
    c10.line.color.rgb = border_gray
    tf10 = c10.text_frame
    tf10.word_wrap = True
    tf10.margin_left = tf10.margin_top = Inches(0.4)

    p = tf10.paragraphs[0]
    p.text = "Principales Aprendizajes y Conclusiones:"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = navy
    p.space_after = Pt(10)

    p = tf10.add_paragraph()
    p.text = (
        "1. Ajuste confiable: La regresión polinómica cúbica ofreció una muy buena aproximación global (R² = 0.9995, Sy/x = 0.927 M) para representar la acumulación de reproducciones a lo largo del año.\n\n"
        "2. Eficiencia comparada de raíces: Newton-Raphson demostró ser 5 veces más rápido que Bisección (3 iteraciones frente a 15) para alcanzar la tolerancia de 10⁻⁴ %, convergiendo a la misma raíz (t* ≈ 35.52 semanas).\n\n"
        "3. Validación del modelo: La discrepancia de 2.75 días respecto a la interpolación lineal directa confirma la coherencia física y estadística del ajuste dentro del error estándar.\n\n"
        "4. Dinámica del consumo: La diferenciación numérica reveló que el pico real de streams ocurrió en la semana 26 (fines de julio de 2023, 4.171 M), y evidenció cómo el modelo cúbico suaviza la respuesta temporal.\n\n"
        "5. Aporte a la ingeniería: El modelado numérico permite convertir datos discretos de telemetría en funciones continuas útiles para dimensionar infraestructura y servidores de streaming ante variaciones de demanda.\n\n"
        "6. Generalización y validación automática: El método resultó reutilizable para otras canciones; demostró que un R² alto (> 0.996) no basta para validar un modelo físico y que el sistema alerta cuándo la cúbica no aplica (estrenos con a > 0)."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = dark_text

    # -------------------------------------------------------------
    # SLIDE 11: Cierre / Preguntas
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    bg11 = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg11.fill.solid()
    bg11.fill.fore_color.rgb = navy
    bg11.line.fill.background()

    tb11 = s11.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.5))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "¡Muchas Gracias por su atención!"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = white
    p.space_after = Pt(20)

    p = tf11.add_paragraph()
    p.text = "¿Preguntas o Consultas?"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(189, 215, 238)
    p.space_after = Pt(30)

    p = tf11.add_paragraph()
    p.text = "Pedro Gatti · Facundo Grillo · Gonzalo Marin\nTaller de Métodos Numéricos · UCC 2026"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(220, 220, 220)

    pptx_path = os.path.join(base_dir, 'Presentacion_Rosa_Pastel_Metodos_Numericos.pptx')
    prs.save(pptx_path)
    print(f"[OK] Presentación PowerPoint guardada exitosamente en: {pptx_path}")


if __name__ == '__main__':
    create_deck()
